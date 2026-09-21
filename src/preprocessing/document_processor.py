import os
import subprocess

from pptx import Presentation
from langchain_community.document_loaders import UnstructuredExcelLoader
from langchain_core.documents import Document


def process_docx(file_path: str) -> str:
    """
    Convert DOCX to PDF.

    Returns:
        Path to the generated PDF.
    """

    output_dir = os.path.dirname(file_path) or "."

    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            output_dir,
            file_path,
        ],
        check=True,
        timeout=60,
    )

    pdf_path = os.path.join(
        output_dir,
        os.path.splitext(
            os.path.basename(file_path)
        )[0] + ".pdf",
    )

    if not os.path.exists(pdf_path):
        raise RuntimeError(
            f"LibreOffice failed to create PDF: {file_path}"
        )

    return pdf_path


def process_pptx(file_path: str) -> list[Document]:
    """
    Extract text, tables, and speaker notes from PPTX.

    One Document is created per slide.
    """

    documents = []

    presentation = Presentation(file_path)

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        text_parts = []

        for shape in slide.shapes:

            if shape.has_text_frame:
                text = shape.text.strip()

                if text:
                    text_parts.append(text)

            if shape.has_table:
                table = shape.table

                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    )

                    if row_text:
                        text_parts.append(row_text)

        # Include speaker notes when available.
        if slide.has_notes_slide:
            notes_frame = (
                slide.notes_slide.notes_text_frame
            )

            if notes_frame:
                notes_text = notes_frame.text.strip()

                if notes_text:
                    text_parts.append(
                        f"Notes: {notes_text}"
                    )

        full_text = "\n".join(
            text_parts
        ).strip()

        if not full_text:
            continue

        documents.append(
            Document(
                page_content=full_text,
                metadata={
                    "source": file_path,
                    "filename": os.path.basename(file_path),
                    "preprocessor": "pptx",
                    "page": slide_number,
                    "slide": slide_number,
                },
            )
        )

    return documents


def process_xlsx(file_path: str) -> list[Document]:
    """
    Extract content from an Excel workbook.

    One or more Documents may be created per sheet/table
    depending on what Unstructured extracts.
    """

    documents = []

    loader = UnstructuredExcelLoader(
        file_path,
        mode="elements",
    )

    extracted_documents = loader.load()

    for doc in extracted_documents:
        text = doc.page_content.strip()

        if not text:
            continue

        sheet_name = doc.metadata.get(
            "page_name",
            "unknown_sheet",
        )

        documents.append(
            Document(
                page_content=text,
                metadata={
                    "source": file_path,
                    "filename": os.path.basename(file_path),
                    "preprocessor": "xlsx",
                    "sheet_name": sheet_name,
                    "page": doc.metadata.get(
                        "page",
                        0,
                    ),
                },
            )
        )

    return documents


def process_document(file_path: str) -> dict:
    """
    Detect the file type and prepare it for ingestion.

    PDF:
        Process later with MinerU.

    DOCX:
        Convert to PDF, then process the PDF with MinerU.

    PPTX:
        Extract directly into LangChain Documents.

    XLSX:
        Extract directly into LangChain Documents.
    """

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":
        return {
            "type": "mineru",
            "path": file_path,
            "documents": None,
            "temporary_pdf": False,
        }

    if extension == ".docx":
        pdf_path = process_docx(file_path)

        return {
            "type": "mineru",
            "path": pdf_path,
            "documents": None,
            "temporary_pdf": True,
        }

    if extension == ".pptx":
        return {
            "type": "documents",
            "path": file_path,
            "documents": process_pptx(file_path),
            "temporary_pdf": False,
        }

    if extension == ".xlsx":
        return {
            "type": "documents",
            "path": file_path,
            "documents": process_xlsx(file_path),
            "temporary_pdf": False,
        }

    raise ValueError(
        f"Unsupported file type: {extension}"
    )